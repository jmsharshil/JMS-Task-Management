"""
AI planning service — turns a SOW/FDD into a day-wise, developer-wise plan
(the same structure as JMS Tech's delivery sheets: D1..Dn, one task per
developer per working day, module-tagged, weekends excluded).
"""
import base64
import json
import re
from django.conf import settings
from openai import AzureOpenAI

MODEL = "gpt-4o-mini"


def _client():
    if not settings.AZURE_OPENAI_API_KEY:
        raise RuntimeError("AZURE_OPENAI_API_KEY is not configured.")
    if not settings.AZURE_OPENAI_ENDPOINT:
        raise RuntimeError("AZURE_OPENAI_ENDPOINT is not configured.")
    return AzureOpenAI(api_key=settings.AZURE_OPENAI_API_KEY,api_version="2024-05-01-preview",azure_endpoint=settings.AZURE_OPENAI_ENDPOINT)


def _ask(content, max_tokens=3000):
    msg = _client().chat.completions.create(model=MODEL, max_completion_tokens=max_tokens,
                                    messages=[{"role": "user", "content": content}])
    return msg.choices[0].message.content


def _parse_json(text):
    clean = re.sub(r"```json|```", "", text).strip()
    start = min([i for i in (clean.find("["), clean.find("{")) if i != -1], default=0)
    return json.loads(clean[start:])


def _format_members(items, default_role="Dev"):
    """Normalize list of User objects or (name, designation) tuples into a '; '-joined string.
    Used to make leader/dev prompts robust regardless of call site (views vs services)."""
    if not items:
        return ""
    first = items[0]
    if hasattr(first, "get_full_name") or hasattr(first, "email"):
        # User model instances
        pairs = [(u.get_full_name() or u.email, getattr(u, "designation", default_role))
                 for u in items]
    else:
        # Already normalized tuples from build_plan_rows
        pairs = items
    return "; ".join(f"{n} ({d})" for n, d in pairs)


def extract_brief(file_bytes=None, filename="", doc_text=""):
    """Condense a SOW/FDD (PDF, DOCX, or pasted text) into {summary, modules[]}."""
    ask = ('Extract a concise scope brief from this SOW/FDD. Respond ONLY with JSON, '
           'no markdown: {"summary":"<max 150 words>","modules":["<short module names '
           "like 'M1 Users', max 12>\"]}")
    if file_bytes:
        import io
        filename = (filename or "").lower()
        if filename.endswith(".docx") or filename.endswith(".doc"):
            try:
                from docx import Document
                document = Document(io.BytesIO(file_bytes))
                doc_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            except Exception as e:
                doc_text = f"[DOCX Extraction Failed: {e}]"
        else:
            from pypdf import PdfReader
            try:
                reader = PdfReader(io.BytesIO(file_bytes))
                doc_text = "\n".join(page.extract_text() or "" for page in reader.pages)
            except Exception as e:
                doc_text = f"[PDF Extraction Failed: {e}]"
            
    content = f"{ask}\n\nDOCUMENT TEXT:\n{doc_text[:12000]}"
    return _parse_json(_ask(content, max_tokens=800))


def weekly_plan(name, weeks, devs, brief, leaders=None, week_from=1, change_note="", done_titles=None):
    """devs: list of (name, designation) tuples. leaders: optional list of same.
    Returns [{week, focus, modules[]}]. Uses architecture context from brief if present.
    Leaders are highlighted for coordination, architecture alignment, reviews, complex modules.
    Backward-compatible with legacy projects that have no architecture in brief."""
    dev_line = _format_members(devs, "Developer")
    leader_line = ""
    if leaders:
        lead_line = _format_members(leaders, "Lead")
        leader_line = f"\nLEADERS (assign coordination, architecture alignment, reviews, complex modules): {lead_line}."

    arch = brief.get("architecture", {})
    prompt = (
        f"Project: {name}. Total {weeks} weeks; plan weeks {week_from}-{weeks}.\n"
        f"Team: {dev_line}{leader_line}\n"
        f"Scope: {brief.get('summary','')}\nModules: {', '.join(brief.get('modules', []))}\n"
    )
    if arch and isinstance(arch, dict):
        prompt += (
            f"\nAPPROVED ARCHITECTURE CONTEXT (must align plan to this):\n"
            f"Overview: {arch.get('overview', '')}\n"
            f"Tech: {arch.get('tech_stack', [])}\n"
            f"Key Components: {arch.get('key_components', [])}\n"
            f"Implementation Approach: {arch.get('implementation_approach', '')}\n"
        )
    if change_note:
        prompt += f"IMPORTANT SCOPE CHANGE (adjust remaining work): {change_note}\n"
    if done_titles:
        prompt += f"Already completed (do not repeat): {'; '.join(done_titles[:30])}\n"
    prompt += (
        "Create a weekly delivery plan that strictly follows the approved architecture "
        "(build first, then integrate, hardening, UAT, launch last). "
        f"Respond ONLY JSON array, no markdown: "
        f'[{{"week":{week_from},"focus":"<max 12 words>","modules":["..."]}}] '
        f"with exactly {weeks - week_from + 1} items for weeks {week_from} to {weeks}."
    )
    return _parse_json(_ask(prompt, max_tokens=1500))


def daily_tasks(name, weeks, week, focus, modules, devs, day_labels, leaders=None, brief=None, change_note=""):
    """One task per developer per working day of the given week.
    devs/leaders: list of (name, designation) tuples. brief may contain approved architecture.
    Leaders get preference for complex/coordination tasks. Architecture context is strictly
    injected and enforced in prompt (aligns with two-stage gated workflow)."""
    dev_line = _format_members(devs, "Developer")
    leader_line = ""
    if leaders:
        lead_line = _format_members(leaders, "Lead")
        leader_line = f"\nLEADERS (prefer for complex tasks, reviews, coordination): {lead_line}."

    arch = brief.get("architecture", {}) if brief and isinstance(brief, dict) else {}
    arch_context = ""
    if arch and isinstance(arch, dict):
        arch_context = (
            f"\nAPPROVED ARCHITECTURE CONTEXT (tasks MUST strictly align to this):\n"
            f"Overview: {arch.get('overview', '')[:150]}\n"
            f"Tech: {arch.get('tech_stack', [])}\n"
            f"Key Components: {arch.get('key_components', [])}\n"
            f"Implementation Approach: {arch.get('implementation_approach', '')[:200]}\n"
        )

    prompt = (
        f'Project "{name}" week {week}/{weeks}. Focus: {focus}. '
        f"Modules this week: {', '.join(modules)}.\n"
        + (f"Scope change in effect: {change_note}\n" if change_note else "")
        + f"Team: {dev_line}{leader_line}{arch_context}\n"
        f"Working days: {', '.join(day_labels)}.\n"
        "Assign exactly ONE task per developer per day, matching each developer's designation. "
        "LEADERS should receive more complex or coordination tasks where possible. "
        "Tasks must align with the approved architecture (components, tech stack, data flows, phased approach). "
        "Tasks concrete, under 12 words. Module names short (e.g. 'M1 Users', 'Setup', 'UAT').\n"
        'Respond ONLY JSON array, no markdown: [{"day":"D6","dev":"<name>","module":"<short>","task":"<short>"}]'
    )
    return _parse_json(_ask(prompt, max_tokens=3000))


def executive_summary(stats_text):
    """On-demand founder summary from live plan stats."""
    return _ask(
        "You are reporting to the founder of JMS Tech. Based on the data below, write a crisp "
        "executive status summary (under 180 words): overall health in one line, what's on track, "
        "what's at risk (name the developers and overdue items), and 2-3 recommended actions. "
        f"Plain text, no markdown headings.\n\n{stats_text}",
        max_tokens=600,
    )


def generate_architecture(project_name: str, team: list, leaders=None, file_bytes=None, filename="", doc_text=""):
    """Generate high-level architecture document from SOW/FDD before detailed planning.
    Returns structured dict with sections for UI rendering (incl. Mermaid diagrams).
    Leaders (subset of team) are explicitly highlighted in the prompt for downstream
    task assignment guidance (coordination, reviews, complex modules). Two-stage
    workflow gate.
    """
    # Extract text if file provided
    if file_bytes:
        import io
        filename = (filename or "").lower()
        if filename.endswith(".docx") or filename.endswith(".doc"):
            try:
                from docx import Document
                document = Document(io.BytesIO(file_bytes))
                doc_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            except Exception as e:
                doc_text = f"[DOCX Extraction Failed: {e}]"
        else:
            from pypdf import PdfReader
            try:
                reader = PdfReader(io.BytesIO(file_bytes))
                doc_text = "\n".join(page.extract_text() or "" for page in reader.pages)
            except Exception as e:
                doc_text = f"[PDF Extraction Failed: {e}]"

    dev_line = _format_members(team, "Developer") if team else "TBD"
    leader_line = ""
    if leaders:
        # For architecture prompt we list names only (different emphasis)
        if leaders and hasattr(leaders[0], "get_full_name"):
            leader_names = [u.get_full_name() or u.email for u in leaders]
        else:
            leader_names = [n for n, _ in leaders]
        leader_line = f"\nTeam Leaders: {', '.join(leader_names)} — assign them coordination, architecture, reviews and complex modules."

    prompt = (
        f"You are a Principal Solutions Architect at JMS Tech. "
        f"Project Name: {project_name}\n"
        f"Team: {dev_line}{leader_line}\n\n"
        f"Based on the following SOW/FDD, produce a comprehensive **Architecture Document**.\n"
        f"Respond **ONLY** with valid JSON (no markdown, no extra text):\n\n"
        "{\n"
        '  "overview": "High-level 2-3 paragraph project overview and architectural approach",\n'
        '  "tech_stack": ["Backend: ...", "Frontend: ...", "Database: ...", "Other tools"],\n'
        '  "key_components": ["Component 1 - description (responsibilities, tech)", ...],\n'
        '  "data_flow": "Detailed description of data flows, integrations, APIs",\n'
        '  "mermaid_diagrams": {\n'
        '    "architecture": "graph TD\\n    Client --> API\\n    ... (valid Mermaid)",\n'
        '    "data_flow": "sequenceDiagram\\n    ..."\n'
        '  },\n'
        '  "non_functional": ["Performance: ...", "Security: ...", "Scalability: ..."],\n'
        '  "risks_assumptions": ["Risk 1 with mitigation", "Assumption 1"],\n'
        '  "implementation_approach": "Phased approach summary"\n'
        "}\n\n"
        f"DOCUMENT TEXT:\n{doc_text[:15000]}"
    )

    try:
        raw = _ask(prompt, max_tokens=4500)
        arch = _parse_json(raw)
        # Ensure all keys exist
        for key in ["overview", "tech_stack", "key_components", "data_flow", "mermaid_diagrams",
                    "non_functional", "risks_assumptions", "implementation_approach"]:
            if key not in arch:
                arch[key] = [] if key in ["tech_stack", "key_components", "non_functional", "risks_assumptions"] else ""
        return arch
    except Exception as e:
        return {"error": str(e), "overview": "Failed to generate architecture. Please try again or provide more details in the SOW."}
